"""EN-08 allow-list, US-401 weekly fold, US-402 test-set parsing and .xlsx, scheduling, Range parsing."""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import Workbook, load_workbook

from chatdocs.adapters.xlsx import testset_xlsx
from chatdocs.application.document_service import RangeNotSatisfiable, parse_range
from chatdocs.application.maintenance import SCHEDULES, slot_for
from chatdocs.domain.analytics.events import SCHEMA, validate_event
from chatdocs.domain.errors import EventInvalid, TestsetInvalid
from chatdocs.domain.evaluation.test_set import parse_rows
from chatdocs.domain.jobs import JobType
from chatdocs.domain.metrics.weekly import fold_weeks, iso_week, week_start
from chatdocs.ports.insights import EvaluationResult

CONTENT_KEYS = ["text", "content", "question", "answer", "excerpt", "prompt", "file_name"]


def test_every_event_in_the_contract_is_listed():
    assert len(SCHEMA) == 13  # dev_tasks.md §0.7 / LG-15


@pytest.mark.parametrize("name", sorted(SCHEMA))
@pytest.mark.parametrize("key", CONTENT_KEYS)
def test_no_event_can_carry_content(name, key):
    """T-EN08.3: whatever the event, a content field is refused (AC-401.3)."""
    with pytest.raises(EventInvalid):
        validate_event(name, {key: "Hợp đồng ABC thanh toán 30%"})


def test_valid_event_and_type_errors():
    assert validate_event("answer_copied", {"with_sources": True}) == {"with_sources": True}
    with pytest.raises(EventInvalid):
        validate_event("answer_copied", {"with_sources": "yes"})
    with pytest.raises(EventInvalid):
        validate_event("citation_clicked", {"citation_index": -1, "doc_removed": False})
    with pytest.raises(EventInvalid):
        validate_event("citation_clicked", {"citation_index": 1})  # missing key
    with pytest.raises(EventInvalid):
        validate_event("document_upload_failed", {"reason": "because"})  # not in the enum
    with pytest.raises(EventInvalid):
        validate_event("made_up", {})
    assert validate_event(
        "answer_completed",
        {
            "latency_ms": 10,
            "first_token_ms": None,
            "citation_count": 0,
            "distinct_docs_cited": 0,
            "is_not_found": True,
            "not_found_reason": "NO_CHUNKS",
        },
    )


def test_weekly_fold():
    today = date(2026, 11, 11)  # Wednesday
    monday = week_start(today)
    daily = {
        monday: {
            "questions": 40,
            "answers": 40,
            "not_found": 4,
            "helpful": 9,
            "feedback_total": 10,
            "p95_latency_ms": 9000,
        },
        monday + timedelta(days=1): {"questions": 60, "answers": 60, "not_found": 6, "p95_latency_ms": 12000},
        monday - timedelta(days=7): {"questions": 10, "answers": 10, "not_found": 5},
    }
    rows = fold_weeks(daily, {monday: 21, monday - timedelta(days=7): 5}, 30, 4, today)
    assert len(rows) == 4 and rows[0].week == iso_week(monday) == "2026-W46"
    w = rows[0]
    assert (w.questions, w.active_users, w.wau_rate) == (100, 21, 0.7)
    assert w.not_found_rate == 0.1 and w.helpful_rate == 0.9 and w.p95_latency_ms == 12000
    assert rows[1].not_found_rate == 0.5 and rows[1].helpful_rate is None
    assert rows[3].questions == 0 and rows[3].wau_rate == 0


def test_test_set_parsing_and_errors():
    header = ["STT", "Câu hỏi", "Đáp án mong đợi", "Không có đáp án (x)"]
    cases = parse_rows([header, [1, "Giá?", "1 tỷ", None], [None, None, None, None], [2, "Màu?", None, "x"]])
    assert [(c.ordinal, c.question, c.expects_not_found) for c in cases] == [(1, "Giá?", False), (2, "Màu?", True)]
    with pytest.raises(TestsetInvalid) as e:
        parse_rows([header, [1, "Giá?"], [2, None, "đáp án"]])
    assert e.value.details == {"row": 3, "column": "Câu hỏi"}
    with pytest.raises(TestsetInvalid) as e:
        parse_rows([["Số", "Hỏi"], [1, "x"]])
    assert e.value.details == {"row": 1, "column": "STT"}
    with pytest.raises(TestsetInvalid) as e:
        parse_rows([header, ["một", "Giá?"]])
    assert e.value.details == {"row": 2, "column": "STT"}
    with pytest.raises(TestsetInvalid):
        parse_rows([header] + [[i, f"Q{i}"] for i in range(1, 202)])  # > 200 cases


def test_xlsx_round_trip(tmp_path):
    path = tmp_path / "t.xlsx"
    path.write_bytes(testset_xlsx.template())
    cases = parse_rows(testset_xlsx.read_rows(str(path)))
    assert len(cases) == 2 and cases[1].expects_not_found
    out = testset_xlsx.build_results(
        [EvaluationResult(cases[0], "Thanh toán 30% [1]", "COMPLETED", [{"index": 1, "file": "a.pdf", "page": 3}], 900)]
    )
    ws = load_workbook(io.BytesIO(out)).active
    assert ws["A1"].value == "STT" and ws["H1"].value == "Chấm: Đúng / Sai"
    assert ws["D2"].value == "Thanh toán 30% [1]" and "a.pdf — trang 3" in ws["F2"].value


def test_xlsx_that_is_not_xlsx(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(TestsetInvalid):
        testset_xlsx.read_rows(str(p))


def test_xlsx_formulas_are_not_evaluated(tmp_path):
    wb = Workbook()
    wb.active.append(["STT", "Câu hỏi"])
    wb.active.append([1, "=1+1"])
    p = tmp_path / "f.xlsx"
    wb.save(p)
    rows = testset_xlsx.read_rows(str(p))
    assert rows[1][1] in (None, "=1+1")  # never 2


def test_schedule_slots():
    sweep = next(s for s in SCHEDULES if s.job_type is JobType.SWEEP_SESSIONS)
    rollup = next(s for s in SCHEDULES if s.job_type is JobType.ROLLUP_METRICS)
    t = datetime(2026, 11, 9, 10, 7, tzinfo=timezone.utc)
    assert slot_for(sweep, t) == slot_for(sweep, t + timedelta(minutes=7))
    assert slot_for(sweep, t) != slot_for(sweep, t + timedelta(minutes=15))
    assert slot_for(rollup, t) is None
    assert slot_for(rollup, t.replace(hour=18)) == slot_for(rollup, t.replace(hour=23)) == "rollup_metrics:2026-11-09"


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        (None, None),
        ("bytes=0-99", (0, 99)),
        ("bytes=100-", (100, 999)),
        ("bytes=-200", (800, 999)),
        ("bytes=900-5000", (900, 999)),
        ("items=0-1", None),
        ("bytes=0-1,5-6", None),
    ],
)
def test_parse_range(header, expected):
    assert parse_range(header, 1000) == expected


def test_parse_range_unsatisfiable():
    with pytest.raises(RangeNotSatisfiable):
        parse_range("bytes=1000-", 1000)
