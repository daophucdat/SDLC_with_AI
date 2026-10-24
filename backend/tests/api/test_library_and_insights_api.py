"""US-204 remove + purge, US-305 file by range, EN-08 events, US-401 metrics, US-402 evaluation, EN-09 headers."""

from __future__ import annotations

import io
from datetime import timedelta

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from chatdocs.adapters.http.main import create_app
from chatdocs.adapters.xlsx import testset_xlsx
from chatdocs.domain.jobs import JobType
from chatdocs.domain.metrics.weekly import week_start
from tests.conftest import csrf_headers, login
from tests.fakes import make_container, seed_ready_document

PAGES = ["Hợp đồng ABC. Thanh toán 30% khi ký hợp đồng.", "Bảo hành 24 tháng."]


async def client_for(f_container, who="member"):
    container, f = f_container
    c = httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app(container)), base_url="https://testserver")
    await login(c, who)
    return c


@pytest.fixture
def stack2():
    return make_container(feature_flags="history,duplicate,evaluation")


# ---------------------------------------------------------------- US-204
async def test_only_uploader_or_lead_can_remove(stack2):
    container, f = stack2
    member = await client_for(stack2, "member")
    other = await client_for(stack2, "member2")
    lead = await client_for(stack2, "lead")
    owner = f.users.by_name("Trần Thu Hà")
    doc = await seed_ready_document(f, "Hop-dong-ABC.pdf", PAGES, uploader=owner)

    listed = (await other.get("/api/v1/documents")).json()["items"][0]
    assert listed["canRemove"] is False  # AC-204.1: the menu item is hidden
    r = await other.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(other))
    assert r.status_code == 403 and r.json()["code"] == "FORBIDDEN"  # AC-204.5
    assert (await lead.get("/api/v1/documents")).json()["items"][0]["canRemove"] is True

    r = await member.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(member))
    assert r.status_code == 202 and r.json() == {"id": str(doc.id), "status": "REMOVED"}
    assert (await member.get("/api/v1/documents")).json()["items"] == []
    assert "remove" in f.audit.actions()
    assert any(j.type is JobType.PURGE_DOCUMENT for j in f.jobs.jobs.values())
    removed = next(r for r in f.events.records if r.name == "document_removed")
    assert removed.props["role"] == "uploader"
    again = await lead.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(lead))
    assert again.status_code == 404
    for c in (member, other, lead):
        await c.aclose()


async def test_lead_removal_is_recorded_as_lead(stack2):
    container, f = stack2
    member = await client_for(stack2, "member")
    lead = await client_for(stack2, "lead")
    doc = await seed_ready_document(f, "a.pdf", PAGES, uploader=f.users.by_name("Trần Thu Hà"))
    assert (await lead.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(lead))).status_code == 202
    assert next(r for r in f.events.records if r.name == "document_removed").props["role"] == "lead"
    await member.aclose()
    await lead.aclose()


async def test_purge_deletes_object_and_chunks_and_keeps_a_tombstone(stack2):
    container, f = stack2
    member = await client_for(stack2)
    doc = await seed_ready_document(f, "a.pdf", PAGES)
    await member.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(member))
    f.clock.advance(hours=21)
    assert await container.documents.purge_overdue() == 1  # T204.3 alert condition
    await container.documents.purge(doc.id)
    assert doc.storage_key not in f.storage.objects
    assert f.index.chunks[doc.id] == [] and f.index.pages[doc.id] == []
    assert doc.id in f.documents.docs and doc.id in f.documents.purged_at  # AC-204.4
    assert await container.documents.purge_overdue() == 0
    await member.aclose()


# ---------------------------------------------------------------- US-305
async def test_file_supports_range_requests(stack2):
    container, f = stack2
    member = await client_for(stack2)
    doc = await seed_ready_document(f, "Hợp đồng ABC.pdf", PAGES)
    data = f.storage.objects[doc.storage_key]

    full = await member.get(f"/api/v1/documents/{doc.id}/file")
    assert full.status_code == 200 and full.content == data
    assert full.headers["accept-ranges"] == "bytes" and full.headers["content-type"] == "application/pdf"
    assert full.headers["cache-control"] == "private, max-age=300"
    assert "filename*=UTF-8''H%E1%BB%A3p" in full.headers["content-disposition"]

    part = await member.get(f"/api/v1/documents/{doc.id}/file", headers={"Range": "bytes=0-99"})
    assert part.status_code == 206 and part.content == data[:100]
    assert part.headers["content-range"] == f"bytes 0-99/{len(data)}" and part.headers["content-length"] == "100"

    bad = await member.get(f"/api/v1/documents/{doc.id}/file", headers={"Range": f"bytes={len(data)}-"})
    assert bad.status_code == 416 and bad.headers["content-range"] == f"bytes */{len(data)}"

    await member.delete(f"/api/v1/documents/{doc.id}", headers=csrf_headers(member))
    gone = await member.get(f"/api/v1/documents/{doc.id}/file")
    assert gone.status_code == 404 and gone.json()["code"] == "DOC_NOT_FOUND"  # EC-16
    await member.aclose()


# ---------------------------------------------------------------- EN-08
async def test_frontend_events(stack2):
    container, f = stack2
    c = await client_for(stack2)
    ok = await c.post(
        "/api/v1/events",
        json={"name": "citation_clicked", "props": {"citation_index": 1, "doc_removed": False}},
        headers=csrf_headers(c),
    )
    assert ok.status_code == 202
    rec = f.events.records[-1]
    assert rec.name == "citation_clicked" and len(rec.user_id_hash) == 64 and rec.app_version == "0.1.0"

    server_only = await c.post(
        "/api/v1/events", json={"name": "answer_completed", "props": {}}, headers=csrf_headers(c)
    )
    assert server_only.status_code == 400 and server_only.json()["code"] == "EVENT_INVALID"
    content = await c.post(
        "/api/v1/events",
        json={"name": "answer_copied", "props": {"with_sources": True, "text": "Hợp đồng ABC"}},
        headers=csrf_headers(c),
    )
    assert content.status_code == 400 and content.json()["field"] == "text"
    await c.aclose()


async def test_event_rate_limit():
    stack = make_container(rate_events_per_min=3)
    c = await client_for(stack)
    body = {"name": "answer_copied", "props": {"with_sources": True}}
    codes = [(await c.post("/api/v1/events", json=body, headers=csrf_headers(c))).status_code for _ in range(4)]
    assert codes == [202, 202, 202, 429]
    await c.aclose()


async def test_login_emits_session_started_and_denial_emits_access_denied(stack2):
    container, f = stack2
    c = await client_for(stack2, "member")
    o = httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app(container)), base_url="https://testserver")
    await login(o, "outsider")
    names = [r.name for r in f.events.records]
    assert names.count("session_started") == 1 and names.count("access_denied") == 1
    assert next(r for r in f.events.records if r.name == "session_started").props == {"is_first_session": True}
    await c.aclose()
    await o.aclose()


# ---------------------------------------------------------------- US-401
async def test_metrics_permissions_and_shape(stack2):
    container, f = stack2
    member = await client_for(stack2, "member")
    lead = await client_for(stack2, "lead")
    manager = await client_for(stack2, "manager")
    monday = week_start(f.clock.now().date())
    f.metrics_repo.days[monday] = {"questions": 50, "answers": 50, "not_found": 5, "helpful": 8, "feedback_total": 10}
    f.metrics_repo.wau[monday] = 21

    assert (await member.get("/api/v1/metrics/weekly")).status_code == 403  # AC-401.2
    for c in (lead, manager):
        r = await c.get("/api/v1/metrics/weekly?weeks=8")
        assert r.status_code == 200 and len(r.json()) == 8
    row = (await lead.get("/api/v1/metrics/weekly")).json()[0]
    assert row == {
        "week": row["week"],
        "weekStart": monday.isoformat(),
        "activeUsers": 21,
        "wauRate": 0.7,
        "questions": 50,
        "notFoundRate": 0.1,
        "helpfulRate": 0.8,
        "p95LatencyMs": None,
    }
    me = (await member.get("/api/v1/me")).json()
    assert "metrics:view" not in me["permissions"] and "evaluation" in me["flags"]
    for c in (member, lead, manager):
        await c.aclose()


async def test_rollup_job_writes_yesterday(stack2):
    container, f = stack2
    yesterday = f.clock.now().date() - timedelta(days=1)
    f.metrics_repo.computed[yesterday] = {"questions": 7.0, "active_users": 3.0}
    await container.metrics.rollup()
    assert f.metrics_repo.days[yesterday] == {"questions": 7.0, "active_users": 3.0}


# ---------------------------------------------------------------- US-402
def xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_evaluation_flow(stack2):
    container, f = stack2
    lead = await client_for(stack2, "lead")
    member = await client_for(stack2, "member")
    await seed_ready_document(f, "Hop-dong-ABC.pdf", PAGES)
    header = ["STT", "Câu hỏi", "Đáp án mong đợi", "Không có đáp án (x)"]
    good = xlsx([header, [1, "Thanh toán hợp đồng ABC?", "30% khi ký", None], [2, "zzqx kkp", None, "x"]])
    files = {"file": ("bo-cau-hoi.xlsx", good, testset_xlsx.__name__)}

    assert (await member.post("/api/v1/evaluations", files=files, headers=csrf_headers(member))).status_code == 403
    r = await lead.post("/api/v1/evaluations", files=files, headers=csrf_headers(lead))
    assert r.status_code == 202 and r.json()["total"] == 2
    run_id = r.json()["runId"]
    assert "evaluation_run" in f.audit.actions()
    assert (await lead.get(f"/api/v1/evaluations/{run_id}/export")).status_code == 409  # not finished

    from uuid import UUID

    assert await container.evaluation.execute(UUID(run_id)) == "done"
    st = (await lead.get(f"/api/v1/evaluations/{run_id}")).json()
    assert st == {"runId": run_id, "status": "done", "done": 2, "total": 2}
    out = await lead.get(f"/api/v1/evaluations/{run_id}/export")
    assert out.status_code == 200
    ws = load_workbook(io.BytesIO(out.content)).active
    assert ws["E2"].value == "COMPLETED" and "Hop-dong-ABC.pdf" in ws["F2"].value
    assert ws["E3"].value == "NOT_FOUND:NO_CHUNKS"

    bad = xlsx([header, [1, None, "x"]])
    r = await lead.post(
        "/api/v1/evaluations", files={"file": ("b.xlsx", bad, "application/octet-stream")}, headers=csrf_headers(lead)
    )
    assert r.status_code == 400 and r.json()["code"] == "TESTSET_INVALID"
    assert (r.json()["row"], r.json()["column"]) == (2, "Câu hỏi")  # AC-402.2
    await lead.aclose()
    await member.aclose()


async def test_evaluation_stops_at_the_deadline_and_keeps_results(stack2):
    container, f = stack2
    lead = await client_for(stack2, "lead")
    await seed_ready_document(f, "Hop-dong-ABC.pdf", PAGES)
    rows = [["STT", "Câu hỏi"], *[[i, f"Thanh toán hợp đồng ABC {i}?"] for i in range(1, 4)]]
    r = await lead.post("/api/v1/evaluations", files={"file": ("t.xlsx", xlsx(rows), "x")}, headers=csrf_headers(lead))
    from uuid import UUID

    run = UUID(r.json()["runId"])
    assert await container.evaluation.execute(run, deadline_s=-1) == "timeout"  # AC-402.3
    assert (await lead.get(f"/api/v1/evaluations/{run}/export")).status_code == 200
    await lead.aclose()


async def test_evaluation_is_hidden_without_the_flag():
    stack = make_container(feature_flags="")
    lead = await client_for(stack, "lead")
    r = await lead.post("/api/v1/evaluations", files={"file": ("t.xlsx", b"x", "x")}, headers=csrf_headers(lead))
    assert r.status_code == 404
    await lead.aclose()


# ---------------------------------------------------------------- EN-09 / EN-07
async def test_security_headers_on_every_response(client):
    for path in ("/healthz", "/api/v1/me"):
        r = await client.get(path)
        assert "frame-ancestors 'none'" in r.headers["content-security-policy"]
        assert r.headers["strict-transport-security"] == "max-age=31536000"
        assert r.headers["x-content-type-options"] == "nosniff"
        assert r.headers["referrer-policy"] == "same-origin"


async def test_prometheus_metrics_have_bounded_labels(stack2):
    container, f = stack2
    c = await client_for(stack2)
    doc = await seed_ready_document(f, "a.pdf", PAGES)
    await c.get(f"/api/v1/documents/{doc.id}/file")
    body = (await c.get("/metrics")).text
    assert "http_requests_total" in body and "/api/v1/documents/{document_id}/file" in body
    assert str(doc.id) not in body
    await c.aclose()


async def test_upload_rate_limit():
    from tests.conftest import pdf_bytes

    stack = make_container(rate_uploads_per_hour=1)
    c = await client_for(stack)
    up = {"file": ("a.pdf", pdf_bytes("Một"), "application/pdf")}
    assert (await c.post("/api/v1/documents", files=up, headers=csrf_headers(c))).status_code == 202
    up = {"file": ("b.pdf", pdf_bytes("Hai"), "application/pdf")}
    r = await c.post("/api/v1/documents", files=up, headers=csrf_headers(c))
    assert r.status_code == 429 and r.json()["code"] == "RATE_LIMITED"
    await c.aclose()


async def test_idle_sessions_are_swept(stack2):
    container, f = stack2
    c = await client_for(stack2)
    assert len(f.sessions.rows) == 1
    f.clock.advance(hours=9)
    assert await container.access.sweep_idle() == 1  # US-103 T103.1
    assert (await c.get("/api/v1/me")).status_code == 401
    await c.aclose()
